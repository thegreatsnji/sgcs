"""Testes Fase 6: categorização, reconciliação e regras de não-fusão (dados fictícios)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from apps.data_migration.current_ops import (
    CAT_DOCUMENTO,
    CAT_MATERIAL,
    CAT_MEDICAMENTO,
    CAT_TESTE,
    PR_IGUAL,
    PR_NOVO_SERVICO,
    PR_PRECO_DIFERENTE,
    PR_AUSENTE_NOVA,
    PR_REVISAR,
    ST_GRAFIA_DIFERENTE,
    ST_NAO_STOCK,
    ST_NOVO_ITEM,
    ST_POSSIVEL_DUPLICADO,
    build_price_comparison,
    build_stock_reconciliation,
    categorize_item,
    compare_prices,
    looks_like_pack_quantity,
    never_auto_merge,
    next_catalog_version,
    parse_pack_quantity,
    preserve_original_quantity,
    reconcile_status,
    write_current_pack,
)


class TestCategorization:
    def test_medicine_material_test_and_card(self):
        assert categorize_item("Ceftriaxona 1g") == CAT_MEDICAMENTO
        assert categorize_item("Seringa 10cc") == CAT_MATERIAL
        assert categorize_item("Teste Widal") == CAT_TESTE
        assert categorize_item("Cartão de Vacina") == CAT_DOCUMENTO


class TestQuantityRules:
    def test_cx50_is_not_converted(self):
        assert parse_pack_quantity("CX/50") is None
        assert parse_pack_quantity("2CX/10") is None
        assert parse_pack_quantity("1T") is None
        assert parse_pack_quantity("4L/7CP") is None
        assert parse_pack_quantity("28") is None
        assert looks_like_pack_quantity("CX/50") is True
        assert preserve_original_quantity("CX/50") == "CX/50"


class TestNoAutoMerge:
    def test_ceftriaxona_spellings_stay_separate(self):
        assert never_auto_merge("CETRIAXONA", "CEFRIAZOMA") is True
        assert never_auto_merge("CETROXONA", "CITROXONA") is True
        status = reconcile_status("Ceftriaxona 1g", "CETRIAXONA", CAT_MEDICAMENTO)
        assert status == ST_GRAFIA_DIFERENTE

    def test_novalgina_is_possible_duplicate_not_merged(self):
        status = reconcile_status("Metamizol/Novalgina inj.", "NOVALGINA", CAT_MEDICAMENTO)
        assert status == ST_POSSIVEL_DUPLICADO


class TestReconciliation:
    def test_new_item_and_non_stock_card(self):
        assert reconcile_status("Adrenalina", None, CAT_MEDICAMENTO) == ST_NOVO_ITEM
        assert reconcile_status("Cartão de Grávida", "CARTAO DE GRAVIDA", CAT_DOCUMENTO) == ST_NAO_STOCK

    def test_build_rows_without_merging(self):
        photos = [{"nome_original": "Ceftriaxona 1g", "tipo": CAT_MEDICAMENTO, "observacao": ""}]
        historical = [{"nome_original": "CETROXONA"}, {"nome_original": "CITROXONA"}]
        rows = build_stock_reconciliation(photos, historical)
        photo_row = next(row for row in rows if row["nome_foto"] == "Ceftriaxona 1g")
        assert photo_row["status_reconciliacao"] == ST_GRAFIA_DIFERENTE
        assert never_auto_merge("Ceftriaxona 1g", photo_row["nome_historico"]) is True


class TestPricesAndVersioning:
    def test_divergent_and_new_prices(self):
        assert compare_prices(3000, 3000) == PR_IGUAL
        assert compare_prices(3000, 4000) == PR_PRECO_DIFERENTE
        assert compare_prices(None, 5000) == PR_NOVO_SERVICO
        assert compare_prices(3000, None) == PR_AUSENTE_NOVA

    def test_v1_1_not_created_without_confirmation(self):
        assert next_catalog_version(confirmed=False) is None
        assert next_catalog_version(confirmed=True) == "SAUVIDA_V1_1"

    def test_price_comparison_marks_new_service(self):
        v1 = [{"codigo": "CONS-CLIN-GER", "nome": "Clínico Geral", "categoria": "CONSULTA", "preco_fcfa": "3000"}]
        photos = [
            {"categoria": "CONSULTA", "servico": "Clínico Geral", "preco_nova_fonte": None},
            {"categoria": "CONSULTA", "servico": "Pediatria", "preco_nova_fonte": None},
        ]
        rows = build_price_comparison(v1, photos)
        pediatria = next(row for row in rows if row["servico"] == "Pediatria")
        assert pediatria["estado"] == PR_NOVO_SERVICO
        geral = next(row for row in rows if row["servico"] == "Clínico Geral")
        assert geral["estado"] == PR_REVISAR


class TestSimpleStockPack:
    def test_nurse_sheet_has_no_pharmacy_columns(self, tmp_path: Path):
        stats = write_current_pack(tmp_path)
        assert stats["catalogo_v1_1"] is None
        wb = load_workbook(tmp_path / "stock_final_validacao_enfermagem.xlsx")
        headers = [cell.value for cell in next(wb["stock_enfermagem"].iter_rows(min_row=1, max_row=1))]
        joined = " ".join(str(h).lower() for h in headers)
        assert "fornecedor" not in joined
        assert "margem" not in joined
        assert "pos" not in joined
        compressa = None
        for row in wb["stock_enfermagem"].iter_rows(min_row=2, values_only=True):
            if row[0] == "Compressa":
                compressa = row
        assert compressa is not None
        assert compressa[3] in (None, "")
        csv_text = (tmp_path / "stock_atual_fonte_fotos.csv").read_text(encoding="utf-8")
        assert "CX/50" in csv_text
        assert stats["itens_fotos"] == 23
