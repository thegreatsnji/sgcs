"""Workbook Excel fictício para testes — nunca usar dados reais."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook


def write_synthetic_workbook(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    consultas = wb.active
    consultas.title = "Consultas"
    consultas.append(["Nome", "Telefone", "Data", "Médico", "Serviço", "Preço", "Desconto", "Líquido", "Residência"])
    consultas.append(["MARIA   DA SILVA", "955111222", date(2024, 3, 1), "Dr. Paulo Mendes", "CONSULTA GERAL", 3000, 0, 3000, "Bissau"])
    consultas.append(["Maria da Silva", "955111222", date(2024, 6, 2), "Dr. Paulo Mendes", "CLINICA GERAL", 3000, 0, 3000, "Bissau"])
    consultas.append(["Ana Costa", "955333444", date(2024, 4, 10), "Dra. Lúcia Tavares", "COONSULTA GERAL", 3000, 500, 2500, ""])
    consultas.append(["", "", date(2024, 4, 11), "", "CONSULTA GERAL", 3000, 0, 3000, ""])
    consultas["A6"] = "TOTAL"
    consultas["F6"] = 9000

    pediatricas = wb.create_sheet("Consultas pediátricas")
    pediatricas.append(["Nome", "Telefone", "Data", "Médico", "Serviço", "Preço"])
    pediatricas.append(["João Fictício", "955000111", date(2024, 5, 5), "Dr. Paulo Mendes", "Consulta Pediátrica", 5000])

    controlos = wb.create_sheet("Controlos")
    controlos.append(["Nome", "Telefone", "Data", "Serviço", "Preço"])
    controlos.append(["Maria da Silva", "955111222", date(2024, 7, 1), "CONTROLE", 2000])

    lab = wb.create_sheet("Laboratório")
    lab.append(["Nome", "Telefone", "Data", "Exame", "Preço", "Médico"])
    lab.append(["Ana Costa", "955333444", date(2024, 4, 12), "Hemograma Completo", 4000, "Dr. Paulo Mendes"])
    lab.append(["Ana Costa", "955333444", date(2024, 4, 12), "Glicemia", 4000, "Dr. Paulo Mendes"])
    lab.append(["Pedro Teste", "", date(2024, 8, 1), "Widal", 4000, ""])
    lab.append(["Pedro Teste", "", date(2024, 8, 1), "Gota Espessa", 2000, ""])
    lab.append(["Pedro Teste", "", date(2024, 8, 1), "Helicobacter pylori", 8000, ""])
    lab["E8"] = "#REF!"

    eco = wb.create_sheet("Ecografias")
    eco.append(["Nome", "Data", "Serviço", "Preço", "Médico"])
    eco.append(["Beatriz Demo", date(2024, 9, 9), "Ecografia Geneco-Obstetricia", 10000, "Dra. Lúcia Tavares"])

    cir = wb.create_sheet("Cirurgias")
    cir.append(["Nome", "Data", "Serviço", "Preço", "Médico"])
    cir.append(["Carlos Exemplo", date(2023, 1, 15), "Hérnia", 200000, "Dr. Paulo Mendes"])

    vendas = wb.create_sheet("VENDAS MEDICA")
    vendas.append(["Produto", "Data", "Preço", "Quantidade"])
    vendas.append(["CETRIAXONA", date(2024, 2, 1), 1500, 1])
    vendas.append(["CEFRIAZOMA", date(2024, 2, 2), 1500, 1])
    vendas.append(["CETROXONA", date(2024, 3, 1), 1500, 1])
    vendas.append(["CITROXONA", date(2024, 3, 2), 1500, 1])
    vendas.append(["CAMA", date(2024, 3, 3), 5000, 1])
    vendas.append(["MÃO DE OBRA", date(2024, 3, 3), 2000, 1])
    vendas.append(["SUTURA DE FERIDA", date(2024, 3, 4), 3000, 1])

    materiais = wb.create_sheet("Materiais clínicos")
    materiais.append(["Material", "Data", "Preço"])
    materiais.append(["Luvas esterilizadas", date(2024, 2, 10), 500])
    materiais.append(["Seringa 5ml", date(2024, 2, 11), 200])

    resumo = wb.create_sheet("Resumos financeiros")
    resumo.append(["Descrição", "Total"])
    resumo.append(["TOTAL MÊS", 123456])
    resumo["B3"] = "=SUM(B2)"

    similar = wb.create_sheet("Consultas extra")
    similar.append(["Nome", "Telefone", "Data", "Serviço", "Preço", "Desconto", "Líquido"])
    similar.append(["Maria da Silveira", "", date(2024, 10, 1), "CONSULTA GERAL", 3000, 0, 3000])
    similar.append(["Ana Costa", "955333444", "13/01/1890", "CONSULTA GERAL", 3000, 0, 9999])

    wb.save(path)
    wb.close()
    return path
