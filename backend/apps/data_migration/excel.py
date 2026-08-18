"""Leitura read-only do Excel histórico e classificação de folhas/colunas."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from apps.data_migration.constants import EXCEL_ERRORS, HEADER_ALIASES, SKIP_ROW_KEYWORDS
from apps.data_migration.dates import parse_date
from apps.data_migration.phones import looks_like_phone
from apps.data_migration.text import collapse_spaces, fold_for_match, looks_like_person_name


@dataclass
class SheetAudit:
    name: str
    sheet_type: str
    rows: int
    empty_rows: int
    header_row: int | None
    headers: list[str]
    mapped_columns: dict[str, str]
    formula_cells: int
    excel_errors: int
    invalid_dates: int
    phone_values: int
    financial_cells: int
    patient_cells: int
    doctor_cells: int
    service_cells: int


@dataclass
class WorkbookAudit:
    path: str
    sheets: list[SheetAudit] = field(default_factory=list)
    total_rows: int = 0
    formula_cells: int = 0
    excel_errors: int = 0
    invalid_dates: int = 0
    phone_values: int = 0
    financial_cells: int = 0
    patient_cells: int = 0
    doctor_cells: int = 0
    service_cells: int = 0


@dataclass
class ExcelRow:
    sheet: str
    sheet_type: str
    row_number: int
    values: dict[str, Any]
    raw: dict[str, Any]
    formulas: list[str]
    excel_errors: list[str]


def classify_sheet(name: str) -> str:
    folded = fold_for_match(name)
    if "resumo" in folded or "financeiro" in folded or "total" in folded or folded.startswith("soma"):
        return "RESUMO_FINANCEIRO"
    if "laborat" in folded or folded.startswith("lab") or "analis" in folded:
        return "LABORATORIO"
    if "ecograf" in folded:
        return "ECOGRAFIA"
    if "cirurg" in folded or "cirug" in folded:
        return "CIRURGIA"
    if "consulta" in folded and ("controlo" in folded or "controle" in folded):
        return "CONSULTA"
    if "controlo" in folded or "controle" in folded:
        return "CONTROLO"
    if folded.startswith("material") or "materiai" in folded:
        return "STOCK_MATERIAL"
    if "venda" in folded or "medicament" in folded or "farmac" in folded:
        return "STOCK_MEDICAMENTO"
    if "consulta" in folded or "pediatric" in folded or "pediátr" in folded:
        return "CONSULTA"
    return "OUTRO"


def event_type_for_sheet(sheet_type: str) -> str:
    mapping = {
        "CONSULTA": "CONSULTA",
        "CONTROLO": "CONTROLO",
        "LABORATORIO": "LABORATORIO",
        "ECOGRAFIA": "ECOGRAFIA",
        "CIRURGIA": "CIRURGIA",
    }
    return mapping.get(sheet_type, "OUTRO")


def event_type_for_row(sheet_type: str, tipo_operacao: str = "", descricao: str = "") -> str:
    """Classifica o acto pela folha e, quando existe, pelo tipo/descrição da linha."""
    blob = fold_for_match(f"{tipo_operacao} {descricao}")
    if sheet_type == "LABORATORIO":
        return "LABORATORIO"
    if sheet_type == "ECOGRAFIA":
        return "ECOGRAFIA"
    if sheet_type == "CIRURGIA":
        return "CIRURGIA"
    if sheet_type == "CONTROLO":
        return "CONTROLO"
    if sheet_type == "CONSULTA":
        if "controlo" in blob or "controle" in blob:
            return "CONTROLO"
        return "CONSULTA"
    if "controlo" in blob or "controle" in blob:
        return "CONTROLO"
    if "consulta" in blob:
        return "CONSULTA"
    if "ecograf" in blob:
        return "ECOGRAFIA"
    if "cirurg" in blob or "cirug" in blob:
        return "CIRURGIA"
    if "analis" in blob or "laborat" in blob:
        return "LABORATORIO"
    return event_type_for_sheet(sheet_type)


def _header_key(cell: Any) -> str:
    return fold_for_match(str(cell or ""))


def _alias_matches(folded: str, alias: str) -> bool:
    alias = fold_for_match(alias)
    if not folded or not alias:
        return False
    if folded == alias:
        return True
    if folded.startswith(alias + " ") or folded.endswith(" " + alias):
        return True
    # Cabeçalhos reais como "Nº Telemovel" / "NOME DE MEDICO".
    if len(alias) >= 7 and alias in folded:
        return True
    return False


def map_headers(headers: list[str]) -> dict[str, str]:
    mapped: dict[str, str] = {}
    used: set[int] = set()
    for field, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(headers):
            if index in used:
                continue
            folded = _header_key(header)
            if not folded:
                continue
            if any(_alias_matches(folded, alias) for alias in aliases):
                mapped[field] = header
                used.add(index)
                break
    return mapped


def _is_excel_error(value: Any) -> bool:
    text = str(value or "").strip().upper()
    return text in EXCEL_ERRORS or text.startswith("#")


def _cell_has_formula(cell) -> bool:
    return bool(getattr(cell, "data_type", None) == "f" or (isinstance(cell.value, str) and str(cell.value).startswith("=")))


def detect_header_row(ws: Worksheet, max_scan: int = 15) -> tuple[int, list[str], dict[str, str]]:
    best: tuple[int, list[str], dict[str, str], int] | None = None
    for row_idx in range(1, min(max_scan, ws.max_row or 1) + 1):
        values = [collapse_spaces(str(cell.value)) if cell.value is not None else "" for cell in ws[row_idx]]
        mapped = map_headers(values)
        score = len(mapped)
        if "nome" in mapped:
            score += 2
        if best is None or score > best[3]:
            best = (row_idx, values, mapped, score)
    if best is None or best[3] == 0:
        return 1, [], {}
    return best[0], best[1], best[2]


class HistoricalWorkbook:
    """Abre o Excel em modo leitura. Nunca grava no ficheiro original."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self._formulas = load_workbook(self.path, read_only=False, data_only=False)
        self._values = load_workbook(self.path, read_only=False, data_only=True)

    def close(self) -> None:
        self._formulas.close()
        self._values.close()

    def __enter__(self) -> HistoricalWorkbook:
        return self

    def __exit__(self, *args) -> None:
        self.close()

    def sheet_names(self) -> list[str]:
        return list(self._formulas.sheetnames)

    def iter_clinical_rows(self) -> Iterator[ExcelRow]:
        for name in self.sheet_names():
            yield from self._iter_sheet(name)

    def _iter_sheet(self, name: str) -> Iterator[ExcelRow]:
        ws_f = self._formulas[name]
        ws_v = self._values[name]
        sheet_type = classify_sheet(name)
        header_row, headers, mapped = detect_header_row(ws_f)
        field_index = {field: headers.index(header) for field, header in mapped.items() if header in headers}

        for row_f, row_v in zip(ws_f.iter_rows(min_row=header_row + 1), ws_v.iter_rows(min_row=header_row + 1)):
            row_number = row_f[0].row if row_f else 0
            formulas: list[str] = []
            errors: list[str] = []
            raw: dict[str, Any] = {}
            values: dict[str, Any] = {}
            empty = True
            for cell_f, cell_v in zip(row_f, row_v):
                val = cell_v.value if cell_v.value is not None else cell_f.value
                if val not in (None, ""):
                    empty = False
                if _cell_has_formula(cell_f):
                    formulas.append(str(cell_f.value))
                if _is_excel_error(val) or _is_excel_error(cell_f.value):
                    errors.append(str(val or cell_f.value))
            if empty:
                continue
            for field, idx in field_index.items():
                cell_v = row_v[idx] if idx < len(row_v) else None
                cell_f = row_f[idx] if idx < len(row_f) else None
                raw[field] = getattr(cell_f, "value", None)
                values[field] = getattr(cell_v, "value", None)
                if values[field] is None:
                    values[field] = raw[field]
            if _is_skip_row(values):
                continue
            yield ExcelRow(
                sheet=name,
                sheet_type=sheet_type,
                row_number=row_number,
                values=values,
                raw=raw,
                formulas=formulas,
                excel_errors=errors,
            )

    def audit(self) -> WorkbookAudit:
        report = WorkbookAudit(path=str(self.path))
        for name in self.sheet_names():
            sheet = self._audit_sheet(name)
            report.sheets.append(sheet)
            report.total_rows += sheet.rows
            report.formula_cells += sheet.formula_cells
            report.excel_errors += sheet.excel_errors
            report.invalid_dates += sheet.invalid_dates
            report.phone_values += sheet.phone_values
            report.financial_cells += sheet.financial_cells
            report.patient_cells += sheet.patient_cells
            report.doctor_cells += sheet.doctor_cells
            report.service_cells += sheet.service_cells
        return report

    def _audit_sheet(self, name: str) -> SheetAudit:
        ws_f = self._formulas[name]
        ws_v = self._values[name]
        header_row, headers, mapped = detect_header_row(ws_f)
        rows = 0
        empty_rows = 0
        formulas = 0
        errors = 0
        invalid_dates = 0
        phones = 0
        financial = 0
        patients = 0
        doctors = 0
        services = 0
        for row_f, row_v in zip(ws_f.iter_rows(min_row=1), ws_v.iter_rows(min_row=1)):
            row_empty = True
            for cell_f, cell_v in zip(row_f, row_v):
                val = cell_v.value if cell_v.value is not None else cell_f.value
                if val not in (None, ""):
                    row_empty = False
                if _cell_has_formula(cell_f):
                    formulas += 1
                if _is_excel_error(val) or _is_excel_error(cell_f.value):
                    errors += 1
            if row_empty:
                empty_rows += 1
            else:
                rows += 1
        for excel_row in self._iter_sheet(name):
            _, iso, marker = parse_date(excel_row.values.get("data"))
            if marker or (excel_row.values.get("data") not in (None, "") and not iso):
                invalid_dates += 1
            if excel_row.values.get("preco") not in (None, "") or excel_row.values.get(
                "valor_liquido"
            ) not in (None, ""):
                financial += 1
            phone = excel_row.values.get("telefone")
            if phone not in (None, "") and looks_like_phone(str(phone)):
                phones += 1
            if looks_like_person_name(str(excel_row.values.get("nome") or "")):
                patients += 1
            if looks_like_person_name(str(excel_row.values.get("medico") or "")):
                doctors += 1
            if collapse_spaces(str(excel_row.values.get("descricao") or "")):
                services += 1
        return SheetAudit(
            name=name,
            sheet_type=classify_sheet(name),
            rows=max(0, rows - (1 if header_row else 0)),
            empty_rows=empty_rows,
            header_row=header_row,
            headers=[h for h in headers if h],
            mapped_columns=mapped,
            formula_cells=formulas,
            excel_errors=errors,
            invalid_dates=invalid_dates,
            phone_values=phones,
            financial_cells=financial,
            patient_cells=patients,
            doctor_cells=doctors,
            service_cells=services,
        )


def _is_skip_row(values: dict[str, Any]) -> bool:
    name = fold_for_match(str(values.get("nome") or ""))
    desc = fold_for_match(str(values.get("descricao") or ""))
    blob = f"{name} {desc}".strip()
    if not blob:
        return False
    return any(blob == key or blob.startswith(f"{key} ") for key in SKIP_ROW_KEYWORDS)
