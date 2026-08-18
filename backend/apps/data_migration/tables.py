"""Leitura de CSV ou Excel de revisão (primeira folha, cabeçalho na linha 1)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from apps.data_migration.review import read_csv


def read_table(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return []
        headers = [str(cell or "").strip() for cell in rows[0]]
        result: list[dict[str, str]] = []
        for raw in rows[1:]:
            if raw is None or all(cell in (None, "") for cell in raw):
                continue
            result.append(
                {
                    headers[idx]: "" if idx >= len(raw) or raw[idx] is None else str(raw[idx]).strip()
                    for idx in range(len(headers))
                    if headers[idx]
                }
            )
        return result
    return []


def read_first_existing(paths: list[Path]) -> list[dict[str, str]]:
    for path in paths:
        rows = read_table(path)
        if rows:
            return rows
    return []
