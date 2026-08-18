"""Pacote Excel privado de validação presencial (não versionar)."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.data_migration.apply import build_import_plan
from apps.data_migration.blocked import classify_blocked_events
from apps.data_migration.review import read_csv
from apps.data_migration.text import fold_for_match


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)

PACK_FILES = {
    "duplicados": "pacientes_duplicados_validacao.xlsx",
    "eventos_bloqueados": "eventos_bloqueados_validacao.xlsx",
    "laboratorio": "laboratorio_historico_validacao.xlsx",
    "medicos": "medicos_historicos_validacao.xlsx",
    "stock": "stock_urgencia_validacao.xlsx",
    "datas": "datas_suspeitas_validacao.xlsx",
}


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], decisions: list[str] | None = None, decision_col: int = 0):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 4, 16), 44)
    if decisions and decision_col:
        formula = '"' + ",".join(decisions) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Escolha uma decisão da lista"
        dv.errorTitle = "Decisão inválida"
        ws.add_data_validation(dv)
        last = max(len(rows) + 1, 2)
        letter = get_column_letter(decision_col)
        dv.add(f"{letter}2:{letter}{last}")
    ws.sheet_properties.tabColor = "1F4E79"
    return ws


def _save(wb: Workbook, path: Path, instructions: str) -> None:
    note = wb.create_sheet("INSTRUCOES", 0)
    note["A1"] = instructions
    note.column_dimensions["A"].width = 80
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def write_validation_pack(staging_dir: Path) -> dict[str, str]:
    pack = staging_dir / "validation_pack"
    pack.mkdir(parents=True, exist_ok=True)
    plan = build_import_plan(staging_dir, skip_blocked=True)
    historico = read_csv(staging_dir / "historico_clinico_sauvida.csv")
    date_decision = plan.get("_date_decision") or {}
    blocked_ids = {row.get("migration_id") for row in plan["_blocked_patients"]}
    classified = classify_blocked_events(
        plan["_blocked_events"],
        blocked_patients=blocked_ids,
        date_decision=date_decision,
    )
    dups = read_csv(staging_dir / "duplicados_para_validacao_clinica.csv") or read_csv(
        staging_dir / "duplicados_pacientes.csv"
    )
    written: dict[str, str] = {}

    wb = Workbook()
    wb.remove(wb.active)
    dup_rows = []
    for pair in dups:
        dup_rows.append(
            [
                pair.get("grupo") or "",
                pair.get("paciente_a_nome") or "",
                pair.get("paciente_b_nome") or "",
                pair.get("paciente_a_telefone") or "",
                pair.get("paciente_b_telefone") or "",
                f"{pair.get('paciente_a_primeiro_registo') or ''} – {pair.get('paciente_a_ultimo_registo') or ''}",
                f"{pair.get('paciente_b_primeiro_registo') or ''} – {pair.get('paciente_b_ultimo_registo') or ''}",
                pair.get("paciente_a_eventos") or "",
                pair.get("paciente_b_eventos") or "",
                pair.get("paciente_a_tipos") or "",
                pair.get("paciente_b_tipos") or "",
                pair.get("decisao") or "",
            ]
        )
    _sheet(
        wb,
        "duplicados",
        [
            "identificador",
            "nome_a",
            "nome_b",
            "telefone_a",
            "telefone_b",
            "periodo_a",
            "periodo_b",
            "n_eventos_a",
            "n_eventos_b",
            "tipos_atendimento_a",
            "tipos_atendimento_b",
            "decisao",
        ],
        dup_rows,
        ["MESMA_PESSOA", "PESSOAS_DIFERENTES", "INDETERMINADO"],
        12,
    )
    path = pack / PACK_FILES["duplicados"]
    _save(wb, path, "INDETERMINADO = não fundir e não importar neste lote. Canónico em MESMA_PESSOA = menor migration_id.")
    written["duplicados"] = str(path)

    wb = Workbook()
    wb.remove(wb.active)
    ev_rows = []
    for event in classified["_eventos"]:
        ev_rows.append(
            [
                event.get("historico_id") or "",
                event.get("motivo_bloqueio") or "",
                event.get("descricao_original") or "",
                event.get("data_original") or event.get("data_evento") or "",
                event.get("migration_patient_id") or "",
                event.get("folha_origem") or "",
                "",
                "",
            ]
        )
    _sheet(
        wb,
        "eventos",
        [
            "migration_event_id",
            "causa",
            "descricao_original",
            "data_original",
            "paciente_ou_grupo",
            "folha",
            "decisao",
            "observacoes",
        ],
        ev_rows,
        ["ACEITAR", "MANTER_BLOQUEADO", "MANTER_TEXTUAL", "INDETERMINADO"],
        7,
    )
    path = pack / PACK_FILES["eventos_bloqueados"]
    _save(wb, path, "Causas: PACIENTE_DUPLICADO, LAB_AMBIGUO, DATA_SUSPEITA, SEM_DATA, OUTRO.")
    written["eventos_bloqueados"] = str(path)

    exames = read_csv(staging_dir / "mapeamento_exames_historicos.csv")
    possiveis: dict[str, list[str]] = defaultdict(list)
    for row in exames:
        key = fold_for_match(row.get("descricao_excel") or row.get("descricao_normalizada") or "")
        code = (row.get("servico_codigo_sgcs") or row.get("exame_codigo_sgcs") or "").strip()
        if key and code and code not in possiveis[key]:
            possiveis[key].append(code)
    by_desc: dict[str, dict] = {}
    for event in historico:
        if event.get("tipo_evento") != "LABORATORIO" or (event.get("estado_mapeamento") or "") != "AMBIGUO":
            continue
        key = fold_for_match(event.get("descricao_original") or "")
        bucket = by_desc.setdefault(
            key,
            {
                "descricao": event.get("descricao_original") or "",
                "sgcs": " | ".join(possiveis.get(key) or ([event.get("servico_codigo_sgcs")] if event.get("servico_codigo_sgcs") else [])),
                "n": 0,
            },
        )
        bucket["n"] += 1
    lab_rows = [[row["descricao"], row["sgcs"], row["n"], "", "", ""] for row in sorted(by_desc.values(), key=lambda item: -item["n"])]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(
        wb,
        "laboratorio",
        ["descricao_historica", "possiveis_exames_sgcs", "n_ocorrencias", "decisao", "exame_confirmado", "observacoes"],
        lab_rows,
        ["CONFIRMAR_MAPEAMENTO", "MANTER_TEXTUAL", "OUTRO_EXAME", "INDETERMINADO"],
        4,
    )
    path = pack / PACK_FILES["laboratorio"]
    _save(wb, path, "MANTER_TEXTUAL é uma decisão válida. Não é obrigatório mapear.")
    written["laboratorio"] = str(path)

    wb = Workbook()
    wb.remove(wb.active)
    med_rows = [
        [row.get("nome_original") or "", row.get("utilizador_sgcs") or "", "", "Não criar utilizador automaticamente."]
        for row in read_csv(staging_dir / "medicos_historicos.csv")
    ]
    _sheet(
        wb,
        "medicos",
        ["nome_historico", "utilizador_sgcs_sugerido", "decisao", "observacoes"],
        med_rows,
        ["MAPEAR_COM_MEDICO_SGCS", "MANTER_NOME_HISTORICO"],
        3,
    )
    path = pack / PACK_FILES["medicos"]
    _save(wb, path, "Não criar contas. MAPEAR só se o médico SGCS já existir.")
    written["medicos"] = str(path)

    wb = Workbook()
    wb.remove(wb.active)
    date_rows = []
    seen = set()
    for hid, row in date_decision.items():
        seen.add(hid)
        date_rows.append(
            [
                hid,
                row.get("data_original") or "",
                row.get("contexto") or row.get("folha") or "",
                row.get("data_interpretada") or "",
                row.get("data_confirmada") or "",
                row.get("decisao") or "",
                row.get("observacoes") or "",
            ]
        )
    for event in classified["_eventos"]:
        if event.get("motivo_bloqueio") != "SEM_DATA":
            continue
        hid = event.get("historico_id") or ""
        if hid in seen:
            continue
        date_rows.append(
            [
                hid,
                event.get("data_original") or "",
                event.get("folha_origem") or "",
                "",
                "",
                "",
                "Sem data interpretável",
            ]
        )
    _sheet(
        wb,
        "datas",
        ["migration_event_id", "data_original", "contexto", "sugestao", "data_confirmada", "decisao", "observacao"],
        date_rows,
        ["CONFIRMAR_DATA", "CORRIGIR_DATA", "NAO_IMPORTAR", "INDETERMINADO"],
        6,
    )
    path = pack / PACK_FILES["datas"]
    _save(wb, path, "Nunca corrigir automaticamente. CORRIGIR_DATA exige data_confirmada preenchida pela clínica.")
    written["datas"] = str(path)

    wb = Workbook()
    wb.remove(wb.active)
    stock_rows = [
        [
            row.get("nome_original") or "",
            row.get("possivel_nome_canonico") or row.get("nome_normalizado") or "",
            row.get("tipo") or "",
            "",
            "",
            row.get("unidade") or "",
            "",
            "",
            "",
            "Urgência apenas. Sem fornecedor, margem, compras, POS ou lote obrigatório.",
        ]
        for row in read_csv(staging_dir / "stock_referencia_sauvida.csv")
    ]
    _sheet(
        wb,
        "stock",
        [
            "nome_observado",
            "nome_sugerido",
            "tipo_sugerido",
            "manter_no_stock",
            "nome_final",
            "unidade",
            "quantidade_actual",
            "stock_minimo",
            "validade_opcional",
            "observacoes",
        ],
        stock_rows,
        ["SIM", "NAO"],
        4,
    )
    path = pack / PACK_FILES["stock"]
    _save(wb, path, "Preencher só o que a urgência precisa. Não importar nesta fase.")
    written["stock"] = str(path)
    return written


def write_review_workbooks(staging_dir: Path) -> dict[str, str]:
    return write_validation_pack(staging_dir)
