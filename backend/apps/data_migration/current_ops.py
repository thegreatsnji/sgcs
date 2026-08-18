"""Reconciliação operacional actual (fotos) vs histórico e catálogo V1.

Não escreve na BD. Não altera SAUVIDA_V1. Não converte CX/50 em unidades.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.data_migration.paths import data_dir
from apps.data_migration.text import fold_for_match

CAT_MEDICAMENTO = "MEDICAMENTO"
CAT_MATERIAL = "MATERIAL_CLINICO"
CAT_TESTE = "TESTE_RAPIDO"
CAT_DOCUMENTO = "DOCUMENTO_CARTAO"
CAT_PROCEDIMENTO = "PROCEDIMENTO"
CAT_OUTRO = "OUTRO"

ST_JA_EXISTE = "JA_EXISTE"
ST_NOVO_ITEM = "NOVO_ITEM"
ST_POSSIVEL_DUPLICADO = "POSSIVEL_DUPLICADO"
ST_GRAFIA_DIFERENTE = "GRAFIA_DIFERENTE"
ST_CONFLITO = "CONFLITO"
ST_NAO_STOCK = "NAO_STOCK"

PR_IGUAL = "IGUAL"
PR_PRECO_DIFERENTE = "PRECO_DIFERENTE"
PR_NOVO_SERVICO = "NOVO_SERVICO"
PR_AUSENTE_NOVA = "SERVICO_AUSENTE_NOVA_FONTE"
PR_NOME_DIFERENTE = "NOME_DIFERENTE"
PR_REVISAR = "REVISAR"

DEC_AGUARDA_ENFERMEIRA = "AGUARDA_ENFERMEIRA"
DEC_AGUARDA_DIRECCAO = "AGUARDA_VALIDACAO_CLINICA"
DEC_NAO_STOCK = "NAO_INCLUIR_STOCK"

PACK_QTY = re.compile(
    r"(?i)\b\d*\s*CX\s*/\s*\d+\b|\b\d+CX\s*/\s*\d+\b|\b\d+T\b|\b\d+L\s*/\s*\d+CP\b"
)

CEFTRIAXONA_FOLDS = {
    fold_for_match(name)
    for name in (
        "CEFTRIAXONA",
        "CETRIAXONA",
        "CEFRIAZOMA",
        "CETROXONA",
        "CITROXONA",
        "CEFRIZOMA",
        "CEFTRIAXONA 1G",
    )
}

MATERIAL_KEYS = (
    "cateter",
    "cateta",
    "seringa",
    "siringa",
    "siriga",
    "compressa",
    "luva",
    "bisturi",
    "besturi",
    "sonda",
    "clamp",
    "adesivo",
    "gaze",
)
TEST_KEYS = ("teste", "widal", "hbsag", "hcg", "malaria", "glicemia", "gravidez", "gravides")
DOC_KEYS = ("cartao", "vacina", "varcina", "gravida")
PROC_KEYS = ("sutura", "mao de obra")


def categorize_item(name: str, hinted: str = "") -> str:
    hinted = (hinted or "").strip().upper()
    if hinted in {
        CAT_MEDICAMENTO,
        CAT_MATERIAL,
        CAT_TESTE,
        CAT_DOCUMENTO,
        CAT_PROCEDIMENTO,
        CAT_OUTRO,
    }:
        return hinted
    folded = fold_for_match(name)
    if any(key in folded for key in DOC_KEYS) and "teste" not in folded:
        return CAT_DOCUMENTO
    if any(key in folded for key in PROC_KEYS):
        return CAT_PROCEDIMENTO
    if any(key in folded for key in TEST_KEYS):
        return CAT_TESTE
    if any(key in folded for key in MATERIAL_KEYS):
        return CAT_MATERIAL
    if folded in {"cama"}:
        return CAT_OUTRO
    return CAT_MEDICAMENTO


def parse_pack_quantity(text: str) -> int | None:
    """Nunca converte CX/50, 2CX/10, 1T, 4L/7CP nem inteiros isolados."""
    del text
    return None


def preserve_original_quantity(text: str) -> str:
    return (text or "").strip()


def looks_like_pack_quantity(text: str) -> bool:
    raw = (text or "").strip()
    if not raw:
        return False
    return bool(PACK_QTY.search(raw) or re.fullmatch(r"\d+", raw))


def suggest_canonical(name: str) -> str:
    folded = fold_for_match(name)
    if folded in CEFTRIAXONA_FOLDS:
        return "CEFTRIAXONA"
    if any(token in folded for token in ("novalgina", "novolgina", "novolvena", "nolotel", "noletel", "nolotil", "metamizol")):
        return "METAMIZOL / NOVALGINA"
    if "diclofenac" in folded or folded.startswith("declo"):
        return "DICLOFENAC"
    if "paracetamol" in folded:
        return "PARACETAMOL"
    if "luva" in folded:
        return "LUVAS"
    if any(token in folded for token in ("siringa", "seringa", "siriga")):
        return "SERINGA"
    if "compressa" in folded:
        return "COMPRESSA"
    if "catet" in folded:
        return "CATETER"
    if any(token in folded for token in ("gravidez", "gravides", "hcg")):
        return "TESTE HCG"
    return (name or "").strip()


def reconcile_status(photo_name: str, historical_name: str | None, category: str) -> str:
    if category in {CAT_DOCUMENTO, CAT_PROCEDIMENTO}:
        return ST_NAO_STOCK
    if not historical_name:
        return ST_NOVO_ITEM
    photo_f = fold_for_match(photo_name)
    hist_f = fold_for_match(historical_name)
    if photo_f == hist_f:
        return ST_JA_EXISTE
    if photo_f in CEFTRIAXONA_FOLDS and hist_f in CEFTRIAXONA_FOLDS:
        return ST_GRAFIA_DIFERENTE
    if suggest_canonical(photo_name) and suggest_canonical(photo_name) == suggest_canonical(historical_name):
        return ST_POSSIVEL_DUPLICADO
    return ST_CONFLITO


def never_auto_merge(left: str, right: str) -> bool:
    return fold_for_match(left) in CEFTRIAXONA_FOLDS and fold_for_match(right) in CEFTRIAXONA_FOLDS


def compare_prices(v1: int | None, nova: int | None) -> str:
    if v1 is None and nova is not None:
        return PR_NOVO_SERVICO
    if nova is None and v1 is not None:
        return PR_AUSENTE_NOVA
    if v1 is None and nova is None:
        return PR_REVISAR
    if v1 == nova:
        return PR_IGUAL
    return PR_PRECO_DIFERENTE


def next_catalog_version(*, confirmed: bool) -> str | None:
    if not confirmed:
        return None
    return "SAUVIDA_V1_1"


def current_ops_dir() -> Path:
    return data_dir() / "private" / "sauvida_atual"


PHOTO_STOCK: list[dict[str, str]] = [
    {"nome_original": "Ceftriaxona 1g", "tipo": CAT_MEDICAMENTO, "apresentacao": "1g", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": "Grafias históricas CETRIAXONA/CEFRIAZOMA/CETROXONA/CITROXONA — não fundir"},
    {"nome_original": "Metamizol/Novalgina inj.", "tipo": CAT_MEDICAMENTO, "apresentacao": "inj.", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": ""},
    {"nome_original": "Diclofenac inj.", "tipo": CAT_MEDICAMENTO, "apresentacao": "inj.", "quantidade_texto_original": "28", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": "Quantidade literal; não interpretar 28 como unidades confirmadas"},
    {"nome_original": "Paracetamol infusão", "tipo": CAT_MEDICAMENTO, "apresentacao": "infusão", "quantidade_texto_original": "4L/7CP", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": "Texto original preservado; sem conversão"},
    {"nome_original": "Dexametazona", "tipo": CAT_MEDICAMENTO, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": ""},
    {"nome_original": "Furosemida", "tipo": CAT_MEDICAMENTO, "apresentacao": "", "quantidade_texto_original": "1T", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": "1T não convertido"},
    {"nome_original": "Adrenalina", "tipo": CAT_MEDICAMENTO, "apresentacao": "", "quantidade_texto_original": "10", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": "10 não convertido em stock físico"},
    {"nome_original": "Oxitocina", "tipo": CAT_MEDICAMENTO, "apresentacao": "", "quantidade_texto_original": "5", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_1", "observacao": ""},
    {"nome_original": "Cateter", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Seringa", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "2CX/10", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": "2CX/10 preservado; não assume 20 unidades"},
    {"nome_original": "Compressa", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "CX/50", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": "CX/50 não interpretado como 50 unidades"},
    {"nome_original": "Luvas", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Bisturi", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Sonda", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Clamp", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Adesivo", "tipo": CAT_MATERIAL, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Teste de Malária", "tipo": CAT_TESTE, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Teste Widal", "tipo": CAT_TESTE, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Teste HBsAg", "tipo": CAT_TESTE, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Teste HCG", "tipo": CAT_TESTE, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Fita de glicemia", "tipo": CAT_TESTE, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": ""},
    {"nome_original": "Cartão de Vacina", "tipo": CAT_DOCUMENTO, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": "Não incluir no stock de medicamentos sem decisão clínica"},
    {"nome_original": "Cartão de Grávida", "tipo": CAT_DOCUMENTO, "apresentacao": "", "quantidade_texto_original": "", "preco_observado_fcfa": "", "fonte_foto": "medicamentos_disponiveis_parte_2", "observacao": "Não incluir no stock de medicamentos sem decisão clínica"},
]

PHOTO_PRICES: list[dict[str, Any]] = [
    {"categoria": "CONSULTA", "servico": "Clínico Geral", "preco_nova_fonte": None},
    {"categoria": "CONSULTA", "servico": "Especialidade", "preco_nova_fonte": None},
    {"categoria": "CONSULTA", "servico": "Pediatria", "preco_nova_fonte": None},
    {"categoria": "CONSULTA", "servico": "Pré-Natal", "preco_nova_fonte": None},
    {"categoria": "CONSULTA", "servico": "Controle", "preco_nova_fonte": None},
    {"categoria": "IMAGIOLOGIA", "servico": "Ecografia Gineco-Obstétrica", "preco_nova_fonte": None},
    {"categoria": "IMAGIOLOGIA", "servico": "Ecografia Morfologia", "preco_nova_fonte": None},
    {"categoria": "IMAGIOLOGIA", "servico": "Ecografia Renal e Abdominal", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Hérnia", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Hidrocele", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Circuncisão", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Abcesso", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Lipoma", "preco_nova_fonte": None},
    {"categoria": "CIRURGIA", "servico": "Prostatectomia", "preco_nova_fonte": None},
    {"categoria": "GINECO_OBSTETRICIA", "servico": "Parto normal", "preco_nova_fonte": None},
    {"categoria": "GINECO_OBSTETRICIA", "servico": "Aplicação/Extração DIU", "preco_nova_fonte": None},
    {"categoria": "GINECO_OBSTETRICIA", "servico": "Cesariana", "preco_nova_fonte": None},
    {"categoria": "GINECO_OBSTETRICIA", "servico": "Gravidez Ectópica", "preco_nova_fonte": None},
    {"categoria": "GINECO_OBSTETRICIA", "servico": "Quisto", "preco_nova_fonte": None},
]

V1_NAME_ALIASES = {
    fold_for_match("Clínico Geral"): "CONS-CLIN-GER",
    fold_for_match("Especialidade"): "CONS-ESP",
    fold_for_match("Pré-Natal"): "CONS-PRE-NATAL",
    fold_for_match("Controle"): "CONS-CONTROLO",
    fold_for_match("Ecografia Gineco-Obstétrica"): "ECO-GO",
    fold_for_match("Ecografia Geneco-Obstetricia"): "ECO-GO",
    fold_for_match("Ecografia Morfologia"): "ECO-MORF",
    fold_for_match("Ecografia Renal e Abdominal"): "ECO-REN-ABD",
    fold_for_match("Hérnia"): "CIR-HERNIA",
    fold_for_match("Parto normal"): "MAT-PARTO-NORMAL",
    fold_for_match("Aplicação/Extração DIU"): "CIR-DIU-APL",
    fold_for_match("Cesariana"): "CIR-CESAR",
    fold_for_match("Gravidez Ectópica"): "CIR-ECTOP",
    fold_for_match("Quisto"): "CIR-QUISTO",
}


def load_v1_catalog(path: Path | None = None) -> list[dict[str, str]]:
    catalog = path or (data_dir() / "releases" / "catalogo_sauvida_v1.csv")
    if not catalog.is_file():
        return []
    with catalog.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def load_historical_stock(path: Path | None = None) -> list[dict[str, str]]:
    base = data_dir() / "private" / "sauvida_migration"
    for candidate in (
        path,
        base / "stock_urgencia_validacao.xlsx",
        base / "stock_validacao_enfermagem.csv",
        base / "stock_referencia_sauvida.csv",
    ):
        if candidate is None or not Path(candidate).is_file():
            continue
        candidate = Path(candidate)
        if candidate.suffix.lower() == ".csv":
            with candidate.open(encoding="utf-8", newline="") as handle:
                return list(csv.DictReader(handle))
    return []


def _hist_name(row: dict[str, str]) -> str:
    return (row.get("nome_original") or row.get("nome_historico") or "").strip()


def _int_price(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(" ", "").replace(",", ".")))
    except ValueError:
        return None


def _best_historical_match(photo_name: str, historical: list[dict[str, str]]) -> dict[str, str] | None:
    photo_f = fold_for_match(photo_name)
    photo_can = fold_for_match(suggest_canonical(photo_name))
    cluster = None
    fuzzy = None
    for row in historical:
        hist_name = _hist_name(row)
        hist_f = fold_for_match(hist_name)
        if hist_f == photo_f:
            return row
        if photo_f in CEFTRIAXONA_FOLDS and hist_f in CEFTRIAXONA_FOLDS:
            cluster = cluster or row
        if photo_can and fold_for_match(suggest_canonical(hist_name)) == photo_can:
            fuzzy = fuzzy or row
    return cluster or fuzzy


def build_stock_reconciliation(photos: list[dict[str, str]], historical: list[dict[str, str]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    matched_hist: set[str] = set()
    for photo in photos:
        name = photo.get("nome_original") or ""
        category = categorize_item(name, photo.get("tipo") or "")
        hist = _best_historical_match(name, historical)
        hist_name = _hist_name(hist) if hist else ""
        if hist_name:
            matched_hist.add(fold_for_match(hist_name))
        status = reconcile_status(name, hist_name or None, category)
        rows.append(
            {
                "nome_historico": hist_name,
                "nome_foto": name,
                "categoria": category,
                "status_reconciliacao": status,
                "nome_canonico_sugerido": suggest_canonical(name),
                "decisao": DEC_NAO_STOCK if status == ST_NAO_STOCK else DEC_AGUARDA_ENFERMEIRA,
                "observacoes": photo.get("observacao") or "",
            }
        )
    for hist in historical:
        hist_name = _hist_name(hist)
        if not hist_name or fold_for_match(hist_name) in matched_hist:
            continue
        category = categorize_item(hist_name, hist.get("tipo") or hist.get("categoria_sugerida") or "")
        status = ST_NAO_STOCK if category in {CAT_DOCUMENTO, CAT_PROCEDIMENTO, CAT_OUTRO} else ST_GRAFIA_DIFERENTE
        rows.append(
            {
                "nome_historico": hist_name,
                "nome_foto": "",
                "categoria": category,
                "status_reconciliacao": status,
                "nome_canonico_sugerido": suggest_canonical(hist_name),
                "decisao": DEC_NAO_STOCK if status == ST_NAO_STOCK else DEC_AGUARDA_ENFERMEIRA,
                "observacoes": "Presente no Excel histórico; não visto nas fotos actuais",
            }
        )
    return rows


def build_price_comparison(v1_rows: list[dict[str, str]], photo_prices: list[dict[str, Any]]) -> list[dict[str, Any]]:
    v1_by_code = {row.get("codigo"): row for row in v1_rows}
    v1_by_name = {fold_for_match(row.get("nome") or ""): row for row in v1_rows}
    matched_codes: set[str] = set()
    out: list[dict[str, Any]] = []
    for item in photo_prices:
        name = item["servico"]
        folded = fold_for_match(name)
        v1 = v1_by_code.get(V1_NAME_ALIASES.get(folded, "")) or v1_by_name.get(folded)
        v1_price = _int_price((v1 or {}).get("preco_fcfa")) if v1 else None
        nova = item.get("preco_nova_fonte")
        if v1 is None:
            estado = PR_NOVO_SERVICO
        elif nova is None:
            estado = PR_REVISAR
        else:
            estado = compare_prices(v1_price, nova)
            if fold_for_match(v1.get("nome") or "") != folded and estado == PR_IGUAL:
                estado = PR_NOME_DIFERENTE
        if v1:
            matched_codes.add(v1.get("codigo") or "")
        out.append(
            {
                "categoria": item["categoria"],
                "servico": name,
                "preco_v1": v1_price if v1_price is not None else "",
                "preco_nova_fonte": nova if nova is not None else "",
                "estado": estado,
                "decisao": DEC_AGUARDA_DIRECCAO if estado in {PR_PRECO_DIFERENTE, PR_NOVO_SERVICO, PR_REVISAR, PR_NOME_DIFERENTE} else "",
                "observacoes": "Preço da foto não assumido automaticamente" if nova is None else "",
            }
        )
    clinical_cats = {"CONSULTA", "CIRURGIA", "ECOGRAFIA", "MATERNIDADE", "PROCEDIMENTO"}
    for row in v1_rows:
        if (row.get("categoria") or "") not in clinical_cats:
            continue
        if (row.get("codigo") or "") in matched_codes:
            continue
        out.append(
            {
                "categoria": row.get("categoria") or "",
                "servico": row.get("nome") or "",
                "preco_v1": row.get("preco_fcfa") or "",
                "preco_nova_fonte": "",
                "estado": PR_AUSENTE_NOVA,
                "decisao": DEC_AGUARDA_DIRECCAO,
                "observacoes": "No V1; não listado no novo preçário fotográfico",
            }
        )
    return out


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], decisions: list[str] | None = None, col: int = 0) -> None:
    ws = wb.create_sheet(title)
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
    for row in rows:
        ws.append(row)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 4, 16), 42)
    if decisions and col:
        formula = '"' + ",".join(decisions) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{max(len(rows) + 1, 2)}")


def write_current_pack(output_dir: Path | None = None) -> dict[str, Any]:
    dest = output_dir or current_ops_dir()
    dest.mkdir(parents=True, exist_ok=True)
    (dest / "fotos").mkdir(exist_ok=True)
    historical = load_historical_stock()
    v1 = load_v1_catalog()
    photos = [{**row, "tipo": categorize_item(row["nome_original"], row["tipo"])} for row in PHOTO_STOCK]
    recon = build_stock_reconciliation(photos, historical)
    prices = build_price_comparison(v1, PHOTO_PRICES)
    divergent = [
        row
        for row in prices
        if row["estado"] in {PR_PRECO_DIFERENTE, PR_NOVO_SERVICO, PR_REVISAR, PR_NOME_DIFERENTE}
    ]

    stock_fields = [
        "nome_original",
        "tipo",
        "apresentacao",
        "quantidade_texto_original",
        "preco_observado_fcfa",
        "fonte_foto",
        "observacao",
    ]
    _write_csv(dest / "stock_atual_fonte_fotos.csv", stock_fields, photos)

    recon_headers = [
        "nome_historico",
        "nome_foto",
        "categoria",
        "status_reconciliacao",
        "nome_canonico_sugerido",
        "decisao",
        "observacoes",
    ]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(
        wb,
        "reconciliacao",
        recon_headers,
        [[row[key] for key in recon_headers] for row in recon],
        [DEC_AGUARDA_ENFERMEIRA, DEC_NAO_STOCK, "CONFIRMAR", "REJEITAR"],
        6,
    )
    wb.save(dest / "reconciliacao_stock.xlsx")
    wb.close()

    nurse_headers = [
        "nome_actual",
        "tipo",
        "unidade",
        "quantidade_actual",
        "stock_minimo",
        "validade_opcional",
        "preco_referencia_fcfa",
        "manter_no_stock",
        "observacoes",
    ]
    nurse_rows = []
    for photo in photos:
        if photo["tipo"] in {CAT_DOCUMENTO, CAT_PROCEDIMENTO}:
            manter, obs = "NAO", "Documento/cartão — não é stock de medicamento"
        else:
            manter, obs = "", "quantidade_actual vazia até a enfermeira confirmar o texto original"
        nurse_rows.append(
            [
                photo["nome_original"],
                photo["tipo"],
                "",
                "",
                "",
                "",
                photo.get("preco_observado_fcfa") or "",
                manter,
                obs,
            ]
        )
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "stock_enfermagem", nurse_headers, nurse_rows, ["SIM", "NAO"], 8)
    note = wb.create_sheet("INSTRUCOES", 0)
    note["A1"] = (
        "Não converter CX/50 em 50 unidades. Preencher quantidade_actual só depois de ver a embalagem. "
        "Sem fornecedor, margem, POS, lote obrigatório."
    )
    wb.save(dest / "stock_final_validacao_enfermagem.xlsx")
    wb.close()

    price_headers = ["categoria", "servico", "preco_v1", "preco_nova_fonte", "estado", "decisao", "observacoes"]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "comparacao", price_headers, [[row[key] for key in price_headers] for row in prices])
    wb.save(dest / "comparacao_precario_atual_v1.xlsx")
    wb.close()

    div_headers = [
        "servico",
        "preco_actual_v1",
        "preco_novo_documento",
        "fonte",
        "decisao",
        "preco_final",
        "confirmado_por",
        "data_confirmacao",
        "observacoes",
    ]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(
        wb,
        "divergentes",
        div_headers,
        [
            [
                row["servico"],
                row["preco_v1"],
                row["preco_nova_fonte"],
                "precario_clinico_foto_3",
                DEC_AGUARDA_DIRECCAO,
                "",
                "",
                "",
                row["observacoes"],
            ]
            for row in divergent
        ],
    )
    wb.save(dest / "precos_divergentes_validacao.xlsx")
    wb.close()

    (dest / "README.txt").write_text(
        "Fonte: transcrição das fotografias operacionais (Medicamentos disponíveis p.1-2 e preçário clínico).\n"
        "As imagens originais devem ficar em fotos/ (gitignored). Não importar para a BD.\n"
        "SAUVIDA_V1 não é alterado. SAUVIDA_V1_1 só após decisão da Direcção.\n",
        encoding="utf-8",
    )

    cats = Counter(row["tipo"] for row in photos)
    statuses = Counter(row["status_reconciliacao"] for row in recon)
    price_states = Counter(str(row["estado"]) for row in prices)
    return {
        "output_dir": str(dest),
        "itens_fotos": len(photos),
        "medicamentos": cats[CAT_MEDICAMENTO],
        "materiais": cats[CAT_MATERIAL],
        "testes_rapidos": cats[CAT_TESTE],
        "documentos_cartoes": cats[CAT_DOCUMENTO],
        "novos_itens": statuses[ST_NOVO_ITEM],
        "ja_conhecidos": statuses[ST_JA_EXISTE],
        "possiveis_duplicados": statuses[ST_POSSIVEL_DUPLICADO],
        "grafias_divergentes": statuses[ST_GRAFIA_DIFERENTE],
        "nao_stock": statuses[ST_NAO_STOCK],
        "precos_divergentes": price_states[PR_PRECO_DIFERENTE],
        "novos_servicos": price_states[PR_NOVO_SERVICO],
        "servicos_iguais": price_states[PR_IGUAL],
        "precos_a_rever": price_states[PR_REVISAR],
        "catalogo_v1_1": next_catalog_version(confirmed=False),
    }
