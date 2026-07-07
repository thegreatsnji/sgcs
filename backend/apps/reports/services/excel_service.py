"""Exportação de relatórios em Excel."""

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font


class ExcelExportService:
    @staticmethod
    def export_report(titulo: str, data: dict[str, Any], clinica_nome: str = "SauVida Clínica") -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Relatório"
        gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M")
        ws.append([clinica_nome])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([titulo])
        ws["A2"].font = Font(bold=True, size=14)
        ws.append([f"Gerado em: {gerado_em}"])
        ws.append([])
        ExcelExportService._write_dict(ws, data)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _write_dict(ws, data: dict, indent: int = 0) -> None:
        for key, value in data.items():
            prefix = "  " * indent
            if isinstance(value, dict):
                ws.append([f"{prefix}{key}"])
                ExcelExportService._write_dict(ws, value, indent + 1)
            elif isinstance(value, list):
                ws.append([f"{prefix}{key}", len(value)])
                for item in value:
                    if isinstance(item, dict):
                        ExcelExportService._write_dict(ws, item, indent + 1)
                    else:
                        ws.append([f"{prefix}-", item])
            else:
                ws.append([f"{prefix}{key}", value])
