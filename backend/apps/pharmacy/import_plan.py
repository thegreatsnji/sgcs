"""Plano de importação do stock de urgência — sem inventar quantidades."""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from apps.pharmacy.constants import CategoriaItemUrgencia
from apps.pharmacy.models import MedicamentoUrgencia

VALIDO = "VALIDO"
PENDENTE_QUANTIDADE = "PENDENTE_QUANTIDADE"
UNIDADE_AUSENTE = "UNIDADE_AUSENTE"
NOME_DUPLICADO = "NOME_DUPLICADO"
QUANTIDADE_INVALIDA = "QUANTIDADE_INVALIDA"
STOCK_MINIMO_INVALIDO = "STOCK_MINIMO_INVALIDO"
VALIDADE_INVALIDA = "VALIDADE_INVALIDA"
NAO_MANTER = "NAO_MANTER"

KEEP_YES = {"SIM", "S", "YES", "1", "TRUE"}
KEEP_NO = {"NAO", "NÃO", "N", "NO", "0", "FALSE"}


def cell(row: dict, *keys: str) -> str:
    for key in keys:
        if key in row and row[key] not in (None, ""):
            return str(row[key]).strip()
    return ""


def parse_int(raw: str) -> int | None:
    text = (raw or "").strip().replace(" ", "").replace(".0", "")
    if not text:
        return None
    if text.startswith("-") and text[1:].isdigit():
        return int(text)
    if not text.isdigit():
        return None
    return int(text)


def parse_date(raw: str) -> date | None:
    text = (raw or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def read_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        from apps.data_migration.review import read_csv

        return read_csv(path)
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        preferred = [name for name in workbook.sheetnames if "stock" in name.lower() or "enferm" in name.lower()]
        names = preferred or workbook.sheetnames
        for name in names:
            sheet = workbook[name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(cell or "").strip() for cell in rows[0]]
            if not any("nome" in header.lower() for header in headers):
                continue
            result: list[dict[str, str]] = []
            for raw in rows[1:]:
                if raw is None or all(cell in (None, "") for cell in raw):
                    continue
                result.append(
                    {
                        headers[idx]: "" if idx >= len(raw) or raw[idx] is None else str(raw[idx]).strip()
                        for idx in range(len(headers))
                        if headers[idx]
                    }
                )
            return result
    finally:
        workbook.close()
    return []


def classify_row(row: dict[str, str], *, seen_names: set[str]) -> dict[str, Any]:
    nome = cell(row, "nome", "nome_actual", "nome_confirmado", "nome_original")
    nome_foto = cell(row, "nome_original", "nome_foto")
    manter = cell(row, "manter_no_stock", "manter_no_stock").upper()
    unidade = cell(row, "unidade")
    qty_raw = cell(row, "quantidade_actual", "quantidade")
    min_raw = cell(row, "stock_minimo", "stock_minimo")
    validade_raw = cell(row, "validade_opcional", "validade")
    tipo = cell(row, "tipo", "categoria") or CategoriaItemUrgencia.MEDICAMENTO
    if tipo not in CategoriaItemUrgencia.values:
        tipo = CategoriaItemUrgencia.OUTRO

    result = {
        "nome": nome,
        "nome_foto": nome_foto,
        "tipo": tipo,
        "unidade": unidade,
        "quantidade_raw": qty_raw,
        "stock_minimo_raw": min_raw,
        "validade_raw": validade_raw,
        "manter": manter,
        "observacoes": cell(row, "observacoes"),
        "quantidade_texto_original": cell(row, "quantidade_texto_original"),
        "quantidade": None,
        "stock_minimo": 5,
        "validade": None,
        "status": VALIDO,
        "ja_existe": False,
    }
    if not nome or manter not in KEEP_YES:
        result["status"] = NAO_MANTER
        return result

    key = nome.casefold()
    if key in seen_names:
        result["status"] = NOME_DUPLICADO
        return result
    seen_names.add(key)

    if MedicamentoUrgencia.objects.filter(nome__iexact=nome).exists():
        result["ja_existe"] = True

    if not unidade:
        result["status"] = UNIDADE_AUSENTE
        return result

    if qty_raw == "":
        result["status"] = PENDENTE_QUANTIDADE
        return result
    qty = parse_int(qty_raw)
    if qty is None or qty < 0:
        result["status"] = QUANTIDADE_INVALIDA
        return result
    result["quantidade"] = qty

    if min_raw:
        minimo = parse_int(min_raw)
        if minimo is None or minimo < 0:
            result["status"] = STOCK_MINIMO_INVALIDO
            return result
        result["stock_minimo"] = minimo

    if validade_raw:
        parsed = parse_date(validade_raw)
        if parsed is None:
            result["status"] = VALIDADE_INVALIDA
            return result
        result["validade"] = parsed

    return result


def build_import_plan(path: Path) -> dict[str, Any]:
    rows = read_rows(path)
    seen: set[str] = set()
    classified = [classify_row(row, seen_names=seen) for row in rows]
    counts = Counter(item["status"] for item in classified)
    valid = [item for item in classified if item["status"] == VALIDO]
    to_create = [item for item in valid if not item["ja_existe"]]
    to_update = [item for item in valid if item["ja_existe"]]
    pending = [item for item in classified if item["status"] == PENDENTE_QUANTIDADE]
    ignored = [item for item in classified if item["status"] != VALIDO]
    return {
        "ficheiro": str(path),
        "encontrados": len(classified),
        "validos": len(valid),
        "a_criar": len(to_create),
        "a_actualizar": len(to_update),
        "ignorados": len(ignored),
        "pendentes": len(pending),
        "por_status": dict(counts),
        "quantidades_iniciais": sum(item["quantidade"] or 0 for item in to_create),
        "itens": classified,
        "a_criar_itens": to_create,
        "a_actualizar_itens": to_update,
    }
